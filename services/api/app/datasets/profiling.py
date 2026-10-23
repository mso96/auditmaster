"""Local parsing and deterministic profiling; no remote readers or services."""

import csv
import io
import math
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pandas as pd
from openpyxl import load_workbook

MAX_BYTES = 50 * 1024 * 1024


def scalar(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip() or None
    return value


def read_table(path: Path, suffix: str, sheet: str | None = None):
    selected = None
    if suffix == ".csv":
        if sheet:
            raise ValueError("CSV dosyasında çalışma sayfası seçilemez.")
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1254")
        if "\x00" in text:
            raise ValueError("Dosya geçerli bir CSV değil.")
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=";,\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            # One-column CSVs are valid; prefer the most common header separator.
            header = text.splitlines()[0] if text else ""
            delimiter = max(";\t,|", key=header.count)
        records = list(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
    else:
        try:
            with ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > 256 * 1024 * 1024:
                    raise ValueError("Excel içeriği açıldığında 256 MB sınırını aşıyor.")
        except BadZipFile as exc:
            raise ValueError("Geçerli bir XLSX dosyası seçin.") from exc
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            selected = sheet or workbook.sheetnames[0]
            if selected not in workbook.sheetnames:
                raise ValueError("Çalışma sayfası bulunamadı.")
            records = list(workbook[selected].iter_rows(values_only=True))
        finally:
            workbook.close()
    records = [list(row) for row in records if any(scalar(v) is not None for v in row)]
    if not records:
        raise ValueError("Dosya boş; sütun başlıkları bulunamadı.")
    headers = [str(v).strip() if v is not None else "" for v in records[0]]
    if not all(headers) or len(headers) != len(set(headers)):
        raise ValueError("Sütun başlıkları dolu ve benzersiz olmalıdır.")
    if any(len(row) != len(headers) for row in records[1:]):
        raise ValueError("Satırlardaki sütun sayıları başlıklarla eşleşmiyor.")
    frame = pd.DataFrame(records[1:], columns=headers, dtype=object)
    return frame, selected


def number(value):
    if isinstance(value, bool):
        raise ValueError
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        text = str(value).strip()
        if re.fullmatch(r"[+-]?\d{1,3}(?:\.\d{3})+(?:,\d+)?", text):
            text = text.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"[+-]?\d+,\d+", text):
            text = text.replace(",", ".")
        elif not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
            raise ValueError
        result = float(text)
    if not math.isfinite(result):
        raise ValueError
    return result


def day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value)
    for pattern in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    raise ValueError


def boolean(value):
    if isinstance(value, bool):
        return value
    text = str(value).casefold()
    if text in ("true", "false", "evet", "hayır"):
        return text in ("true", "evet")
    raise ValueError


def profile(frame: pd.DataFrame):
    columns = []
    for name in frame.columns:
        raw = frame[name].tolist()
        values = [v for v in raw if scalar(v) is not None]
        kind, converted = "text", [str(scalar(v)) for v in values]
        if values:
            for candidate, parser in [("boolean", boolean), ("number", number), ("date", day)]:
                try:
                    converted = [parser(v) for v in values]
                    kind = candidate
                    break
                except (ValueError, TypeError, OverflowError):
                    continue
        counts = Counter(converted)
        columns.append(
            {
                "name": name,
                "type": kind,
                "null_count": len(raw) - len(values),
                "distinct_count": len(counts),
                "min": min(converted) if converted and kind in ("number", "date") else None,
                "max": max(converted) if converted and kind in ("number", "date") else None,
                "top_values": [{"value": v, "count": c} for v, c in counts.most_common(5)]
                if kind == "text"
                else [],
            }
        )
    return {"row_count": len(frame), "columns": columns}


def preview(frame: pd.DataFrame):
    return {
        "columns": list(frame.columns),
        "rows": [
            {key: scalar(value) for key, value in row.items()}
            for row in frame.head(100).to_dict(orient="records")
        ],
    }
