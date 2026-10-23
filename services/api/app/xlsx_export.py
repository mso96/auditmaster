"""Excel export of a saved risk control matrix. openpyxl is pure Python."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.models import RcmRecord

_LEVELS = {"low": "Düşük", "medium": "Orta", "high": "Yüksek"}
_CONTROL_TYPES = {"preventive": "Önleyici", "detective": "Saptayıcı"}
_NATURES = {"manual": "Manuel", "automated": "Otomatik"}
_DECISIONS = {
    "pending": "Bekliyor",
    "accepted": "Kabul",
    "edited": "Düzenlendi",
    "rejected": "Reddedildi",
}
_HEADERS = [
    "Denetim",
    "Kapsam",
    "Risk",
    "Kategori",
    "Olasılık",
    "Etki",
    "Karar",
    "Risk kaynağı",
    "Kontrol",
    "Kontrol tipi",
    "Kontrol doğası",
    "Kontrol kaynağı",
    "Test",
    "Test kaynağı",
]


def build_workbook(audit_title: str, record: RcmRecord) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RCM"
    note = (
        "Öneridir. Nihai karar denetçiye aittir. "
        "Kaynaksız satırlar dayanak pasaj içermez."
    )
    sheet.append([note])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEADERS))
    sheet["A1"].font = Font(italic=True, color="555555")
    sheet.append(_HEADERS)
    for cell in sheet[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for risk in record.risks:
        controls = risk.controls or [None]
        for control in controls:
            tests = risk.tests or [None]
            for test in tests:
                sheet.append(
                    [
                        audit_title,
                        record.scope,
                        risk.description,
                        risk.category,
                        _LEVELS.get(risk.likelihood, risk.likelihood),
                        _LEVELS.get(risk.impact, risk.impact),
                        _DECISIONS.get(risk.decision, risk.decision),
                        _source_cell(risk.citation_status, risk.citations),
                        control.description if control else "",
                        _CONTROL_TYPES.get(control.control_type, "") if control else "",
                        _NATURES.get(control.nature, "") if control else "",
                        _source_cell(control.citation_status, control.citations) if control else "",
                        test.description if test else "",
                        _source_cell(test.citation_status, test.citations) if test else "",
                    ]
                )

    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [28, 32, 42, 22, 12, 12, 14, 36, 42, 16, 16, 36, 42, 36]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A3"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _source_cell(status: str, citations: list) -> str:
    if status == "kaynaksız" or not citations:
        return "Kaynaksız"
    parts: list[str] = []
    for citation in citations:
        where = citation.document_name
        if citation.page is not None:
            where += f", s. {citation.page}"
        if citation.section:
            where += f", {citation.section}"
        parts.append(f"{where}: {citation.snippet}")
    return "\n".join(parts)
