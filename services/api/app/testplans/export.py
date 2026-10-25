from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.testplans.models import PlanRecord

METHODS = {
    "inspection": "İnceleme",
    "observation": "Gözlem",
    "reperformance": "Yeniden uygulama",
    "analytical": "Analitik",
}
DECISIONS = {
    "pending": "Bekliyor",
    "accepted": "Kabul",
    "edited": "Düzenlendi",
    "rejected": "Reddedildi",
}


def export_plan(title: str, plan: PlanRecord) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Test Planı"
    sheet.append(
        ["Öneridir; hiçbir test veya analiz çalıştırılmamıştır. Nihai karar denetçiye aittir."]
    )
    sheet.append(
        [
            "Denetim",
            "Risk",
            "Kontrol",
            "Adım",
            "Test amacı",
            "Yöntem",
            "Prosedür",
            "Örneklem önerisi",
            "İstenecek kanıtlar",
            "İlgili analiz kuralı",
            "Karar",
            "Kaynak durumu",
            "Kaynak atıfları",
        ]
    )
    for index, step in enumerate(plan.steps, 1):
        sources = "\n\n".join(
            f"{c.document_name} | Sayfa: {c.page or '—'} | Bölüm: {c.section or '—'}\n"
            f"{c.snippet}\nBelge: {c.document_id} | Pasaj: {c.chunk_id}"
            for c in step.citations
        )
        sheet.append(
            [
                title,
                step.risk_description,
                step.control_description,
                index,
                step.objective,
                METHODS[step.method],
                step.procedure,
                step.sample_size,
                "\n".join(step.evidence_requests),
                step.related_rule or "Yok",
                DECISIONS[step.decision],
                "Kaynaklı" if step.citations else "Kaynaksız",
                sources,
            ]
        )
    for row in sheet:
        for cell in row:
            # All user/model/library text is text, including leading '=' formulas.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    for column in range(1, 14):
        sheet.column_dimensions[get_column_letter(column)].width = 38 if column != 4 else 8
    sheet.freeze_panes = "E3"
    sheet.auto_filter.ref = f"A2:M{sheet.max_row}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
