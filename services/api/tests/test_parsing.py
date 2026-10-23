import io
from pathlib import Path

import pytest
from app.chunking import chunk_blocks
from app.parsing import ParseError, parse_document
from docx import Document

FIXTURES = Path(__file__).parent / "fixtures"


def test_markdown_fixture_keeps_sections() -> None:
    data = (FIXTURES / "ornek-satin-alma-proseduru.md").read_bytes()
    blocks = parse_document("ornek-satin-alma-proseduru.md", data)
    sections = [block.section for block in blocks]

    assert "Onay yetkisi" in sections
    assert "Teklif alma" in sections
    onay = next(block for block in blocks if block.section == "Onay yetkisi")
    assert "onaylanır" in onay.text
    assert onay.page is None


def test_plain_text_form_feed_sets_pages() -> None:
    text = "Sayfa bir metni burada.\fİkinci sayfa metni burada."
    blocks = parse_document("notlar.txt", text.encode())
    assert [block.page for block in blocks] == [1, 2]
    assert "İkinci sayfa" in blocks[1].text


def test_docx_heading_and_table_keep_order() -> None:
    document = Document()
    document.add_heading("Onay yetkisi", level=1)
    document.add_paragraph("Satın alma talepleri onaylanmadan sipariş açılamaz.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Eşik"
    table.rows[0].cells[1].text = "100.000 TL"
    buffer = io.BytesIO()
    document.save(buffer)

    blocks = parse_document("prosedur.docx", buffer.getvalue())
    assert len(blocks) == 1
    assert blocks[0].section == "Onay yetkisi"
    assert "sipariş açılamaz" in blocks[0].text
    assert "100.000 TL" in blocks[0].text


def test_pdf_pages_are_extracted() -> None:
    payload = _minimal_pdf([["Satin alma onayi zorunludur."], ["Uc yazili teklif aranir."]])
    blocks = parse_document("ornek.pdf", payload)
    assert [block.page for block in blocks] == [1, 2]
    assert "onayi" in blocks[0].text
    assert "teklif" in blocks[1].text


def test_unsupported_extension_and_empty_text() -> None:
    with pytest.raises(ParseError, match="PDF, DOCX"):
        parse_document("tablo.csv", b"a,b")
    with pytest.raises(ParseError, match="metin"):
        parse_document("bos.txt", b"   \n")


def test_chunk_preserves_page_and_overlaps() -> None:
    from app.parsing import ParsedBlock

    page = ParsedBlock(text="Kısa sayfa metni yeterince uzun.", page=3, section="Giriş")
    chunks = chunk_blocks([page])
    assert len(chunks) == 1
    assert chunks[0].page == 3
    assert chunks[0].section == "Giriş"

    words = " ".join(f"kelime{i:04d}" for i in range(400))
    parts = chunk_blocks([ParsedBlock(text=words, page=1, section=None)], size=200, overlap=40)
    assert len(parts) > 1
    assert parts[0].content[-30:] in (parts[0].content + " " + parts[1].content)
    assert parts[1].page == 1


def _minimal_pdf(pages: list[list[str]]) -> bytes:
    objects: dict[int, bytes] = {}
    next_id = 4
    page_ids: list[int] = []
    for lines in pages:
        commands = ["BT", "/F1 12 Tf", "72 740 Td"]
        for index, line in enumerate(lines):
            if index:
                commands.append("0 -16 Td")
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({escaped}) Tj")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1")
        content_id = next_id
        next_id += 1
        page_id = next_id
        next_id += 1
        objects[content_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream"
        )
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 3 0 R >> >> >>"
        ).encode("latin-1")
        page_ids.append(page_id)
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    objects[2] = f"<< /Type /Pages /Count {len(page_ids)} /Kids [{kids}] >>".encode("latin-1")
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    max_id = next_id - 1
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_id in range(1, max_id + 1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode("latin-1"))
        output.extend(objects[object_id])
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {max_id + 1}\n".encode("latin-1"))
    output.extend(b"0000000000 65535 f \n")
    for object_id in range(1, max_id + 1):
        output.extend(f"{offsets[object_id]:010d} 00000 n \n".encode("latin-1"))
    output.extend(
        f"trailer\n<< /Size {max_id + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode(
            "latin-1"
        )
    )
    return bytes(output)
