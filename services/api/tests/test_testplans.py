import io
import json
from pathlib import Path
from uuid import uuid4

import pytest
from app.config import Settings
from app.llm import LLMClient
from app.models import ControlDraft, PassageHit, RiskDraft
from app.testplans.generator import PlanGenerationError, generate_steps
from app.testplans.models import StepEdit
from openpyxl import load_workbook

FIXTURES = Path(__file__).parent / "fixtures"


def control_and_risk():
    control = ControlDraft(
        id=uuid4(),
        description="Sipariş yetkili kişi tarafından onaylanır.",
        control_type="preventive",
        nature="manual",
    )
    risk = RiskDraft(
        id=uuid4(),
        description="Yetkisiz sipariş oluşturulması.",
        category="Onay",
        likelihood="medium",
        impact="high",
        controls=[control],
        decision="accepted",
    )
    return control, risk


def sample_step(**overrides):
    return {
        "objective": "Onay kontrolünün uygulandığını doğrulamak.",
        "method": "inspection",
        "procedure": "Onay kaydını yetki matrisiyle karşılaştırın.",
        "sample_size": "Evren ve riske göre denetçi belirlemelidir.",
        "evidence_requests": ["Yetki matrisi", "Onay kayıtları"],
        "related_rule": "segregation_of_duties",
        "citations": [],
        **overrides,
    }


def test_fake_no_key_no_sdk_and_no_rule_execution(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("No remote model or analytics execution is allowed")

    monkeypatch.setattr("app.llm.OpenAI", forbidden)
    monkeypatch.setattr("app.analytics.engine.run_rules", forbidden)
    control, risk = control_and_risk()
    llm = LLMClient(Settings(fake_models=True, openai_api_key=""))
    steps = generate_steps(risk=risk, control=control, hits=[], llm=llm)
    assert len(steps) == 2
    assert all(s.citation_status == "kaynaksız" and s.decision == "pending" for s in steps)
    assert all(s.control_id == control.id and s.risk_id == risk.id for s in steps)
    assert all(s.evidence_requests and s.sample_size for s in steps)
    again = generate_steps(risk=risk, control=control, hits=[], llm=llm)
    assert [s.model_dump(exclude={"id"}) for s in steps] == [
        s.model_dump(exclude={"id"}) for s in again
    ]


class Stub:
    def __init__(self, responses):
        self.responses = iter(responses)
        self.calls = []

    def complete(self, prompt, **kwargs):
        self.calls.append((prompt, kwargs))
        return next(self.responses)


def test_real_client_schema_repair_and_citation_allowlist():
    control, risk = control_and_risk()
    hit = PassageHit(
        chunk_id=uuid4(),
        document_id=uuid4(),
        filename="Yetki.md",
        doc_type="prosedur",
        page=2,
        section="Onay",
        content="Sipariş onayları yetki matrisiyle karşılaştırılır.",
        score=1,
    )
    payload = {
        "steps": [
            sample_step(
                citations=[
                    {"chunk_id": str(hit.chunk_id), "snippet": hit.content},
                    {"chunk_id": str(uuid4()), "snippet": hit.content},
                    {
                        "chunk_id": str(hit.chunk_id),
                        "snippet": "Bu cümle kaynakta bulunmamaktadır.",
                    },
                ]
            )
        ]
    }
    llm = Stub(["not json", json.dumps(payload)])
    result = generate_steps(risk=risk, control=control, hits=[hit], llm=llm)
    assert len(llm.calls) == 2
    assert llm.calls[0][1]["json_mode"] is True
    assert result[0].citations[0].document_name == "Yetki.md"
    assert len(result[0].citations) == 1
    assert result[0].citation_status == "kaynaklı"
    assert result[0].related_rule == "segregation_of_duties"
    assert result[0].decision == "pending"
    assert str(hit.chunk_id) in llm.calls[0][0]


@pytest.mark.parametrize(
    "patch",
    [
        {"method": "execute"},
        {"related_rule": "run_sql"},
        {"procedure": " "},
        {"evidence_requests": []},
        {"sample_size": "x" * 2001},
    ],
)
def test_invalid_model_payload_rejected_after_one_repair(patch):
    control, risk = control_and_risk()
    raw = json.dumps({"steps": [sample_step(**patch)]})
    llm = Stub([raw, raw])
    with pytest.raises(PlanGenerationError):
        generate_steps(risk=risk, control=control, hits=[], llm=llm)
    assert len(llm.calls) == 2


def seed(client, *, with_source=True, decision="accepted"):
    document = None
    if with_source:
        response = client.post(
            "/documents",
            data={"doc_type": "prosedur", "audit_area": "Satın Alma"},
            files={
                "file": (
                    "ornek-satin-alma-proseduru.md",
                    (FIXTURES / "ornek-satin-alma-proseduru.md").read_bytes(),
                    "text/markdown",
                )
            },
        )
        assert response.status_code == 201, response.text
        document = response.json()
    generated = client.post(
        "/rcm/generate",
        json={
            "scope": "Satın alma onay yetkisi ve teklif kontrolü",
            "audit_area": "Satın Alma",
        },
    )
    assert generated.status_code == 200, generated.text
    body = generated.json()
    for risk in body["risks"]:
        risk["decision"] = decision
    saved = client.post(
        "/rcms", json={key: value for key, value in body.items() if key != "model_mode"}
    )
    assert saved.status_code == 200, saved.text
    return saved.json(), document


def edits(plan):
    return {
        "expected_revision": plan["revision"],
        "steps": [
            StepEdit.model_validate(
                {key: value for key, value in step.items() if key in StepEdit.model_fields}
            ).model_dump(mode="json")
            for step in plan["steps"]
        ],
    }


def test_full_plan_flow_reviews_persistence_sources_export_and_conflicts(client, monkeypatch):
    saved, document = seed(client)
    audit_id = saved["audit_id"]
    url = f"/audits/{audit_id}/test-plan"
    assert client.get(url).json() is None
    assert client.get(url + "/export").status_code == 404
    # No reranker HTTP request even when a deployment has it configured.
    monkeypatch.setenv("RERANKER_URL", "http://must-not-be-called.invalid/rerank")
    from app.config import get_settings

    get_settings.cache_clear()
    response = client.post(url + "/generate", json={"expected_revision": 0})
    assert response.status_code == 200, response.text
    plan = response.json()
    controls = {(r["id"], c["id"]) for r in saved["rcm"]["risks"] for c in r["controls"]}
    assert {(s["risk_id"], s["control_id"]) for s in plan["steps"]} == controls
    assert len(plan["steps"]) == 2 * len(controls)
    assert all(s["decision"] == "pending" for s in plan["steps"])
    assert any(s["citations"] for s in plan["steps"])
    assert any(s["citation_status"] == "kaynaksız" for s in plan["steps"])
    for step in plan["steps"]:
        for citation in step["citations"]:
            assert citation["document_id"] == document["id"]
            assert citation["snippet"] in citation["passage"]
    assert client.post(url + "/generate", json={"expected_revision": 0}).status_code == 409

    body = edits(plan)
    body["steps"][0]["decision"] = "accepted"
    body["steps"][1]["decision"] = "rejected"
    if len(body["steps"]) > 2:
        body["steps"][2]["procedure"] = "=HYPERLINK(1)"
        body["steps"][2]["decision"] = "edited"
    response = client.put(url, json=body)
    assert response.status_code == 200, response.text
    updated = response.json()
    assert updated["revision"] == 2
    assert updated["steps"][0]["decision"] == "accepted"
    assert updated["steps"][1]["decision"] == "rejected"
    if len(updated["steps"]) > 2:
        assert updated["steps"][2]["citation_status"] == "kaynaksız"
        assert updated["steps"][2]["citations"] == []
    assert client.get(url).json() == updated
    assert client.put(url, json=body).status_code == 409
    export = client.get(url + "/export")
    assert export.status_code == 200
    sheet = load_workbook(io.BytesIO(export.content)).active
    assert sheet.title == "Test Planı"
    assert sheet.max_row == len(updated["steps"]) + 2
    assert sheet["K3"].value == "Kabul"
    assert sheet["K4"].value == "Reddedildi"
    assert "ornek-satin-alma-proseduru.md" in sheet["M3"].value
    assert all(c.data_type != "f" for row in sheet for c in row)

    assert client.delete(f"/documents/{document['id']}").status_code == 204
    assert all(s["citation_status"] == "kaynaksız" for s in client.get(url).json()["steps"])
    assert (
        "Kaynaksız"
        == load_workbook(io.BytesIO(client.get(url + "/export").content)).active["L3"].value
    )


@pytest.mark.parametrize("decision", ["pending", "edited", "rejected"])
def test_only_explicitly_accepted_rcm_is_eligible(client, decision):
    saved, _ = seed(client, with_source=False, decision=decision)
    response = client.post(f"/audits/{saved['audit_id']}/test-plan/generate", json={})
    assert response.status_code == 422
    assert "kabul" in response.json()["detail"]


def test_missing_audit_and_untrusted_fields_are_rejected(client):
    assert client.get(f"/audits/{uuid4()}/test-plan").status_code == 404
    saved, _ = seed(client, with_source=False)
    url = f"/audits/{saved['audit_id']}/test-plan"
    plan = client.post(url + "/generate", json={}).json()
    assert all(s["citation_status"] == "kaynaksız" for s in plan["steps"])
    body = edits(plan)
    body["steps"][0]["citations"] = [{"snippet": "invented"}]
    assert client.put(url, json=body).status_code == 422
    body = edits(plan)
    body["steps"][0]["id"] = str(uuid4())
    assert client.put(url, json=body).status_code == 422
    body = edits(plan)
    body["steps"].append(body["steps"][0])
    assert client.put(url, json=body).status_code == 422
    assert client.get(url).json()["revision"] == 1


def test_rcm_changes_invalidate_plan_save_and_export(client):
    saved, _ = seed(client, with_source=False)
    url = f"/audits/{saved['audit_id']}/test-plan"
    plan = client.post(url + "/generate", json={}).json()
    rcm = saved["rcm"]
    rcm["risks"][0]["controls"][0]["description"] = "Denetçi tarafından değişen kontrol."
    response = client.post(
        "/rcms",
        json={
            "audit_id": saved["audit_id"],
            "scope": rcm["scope"],
            "audit_area": rcm["audit_area"],
            "risks": rcm["risks"],
        },
    )
    assert response.status_code == 200
    assert client.get(url).json()["stale"]
    assert client.put(url, json=edits(plan)).status_code == 409
    assert client.get(url + "/export").status_code == 409
    fresh = client.post(url + "/generate", json={"expected_revision": 1})
    assert fresh.status_code == 200
    assert not fresh.json()["stale"]
    assert fresh.json()["steps"][0]["control_description"] == "Denetçi tarafından değişen kontrol."


def test_invalid_generation_does_not_overwrite_existing_plan(client, monkeypatch):
    saved, _ = seed(client, with_source=False)
    url = f"/audits/{saved['audit_id']}/test-plan"
    plan = client.post(url + "/generate", json={}).json()

    def fail(**kwargs):
        raise PlanGenerationError("Geçersiz model yanıtı.")

    monkeypatch.setattr("app.testplans.router.generate_steps", fail)
    assert client.post(url + "/generate", json={"expected_revision": 1}).status_code == 422
    assert client.get(url).json() == plan


def test_cors_allows_save_method(client):
    response = client.options(
        "/audits/123/test-plan",
        headers={
            "Origin": "http://localhost:3000",
            "Access-Control-Request-Method": "PUT",
        },
    )
    assert response.status_code == 200
    assert "PUT" in response.headers["access-control-allow-methods"]


def test_mixed_rcm_generates_only_accepted_controls(client, monkeypatch):
    saved, _ = seed(client)
    rcm = saved["rcm"]
    assert len(rcm["risks"]) >= 2
    rcm["risks"][1]["decision"] = "rejected"
    client.post(
        "/rcms",
        json={
            "audit_id": saved["audit_id"],
            "scope": rcm["scope"],
            "audit_area": rcm["audit_area"],
            "risks": rcm["risks"],
        },
    ).raise_for_status()
    from app.testplans import router as module

    calls = []
    original = module.retrieve

    def tracked(*args, **kwargs):
        calls.append(kwargs)
        return original(*args, **kwargs)

    monkeypatch.setattr(module, "retrieve", tracked)
    plan = client.post(f"/audits/{saved['audit_id']}/test-plan/generate", json={}).json()
    expected = {
        (r["id"], c["id"])
        for r in rcm["risks"]
        if r["decision"] == "accepted"
        for c in r["controls"]
    }
    assert {(s["risk_id"], s["control_id"]) for s in plan["steps"]} == expected
    assert len(calls) == len(expected)
    assert all(c["audit_area"] == "Satın Alma" for c in calls)
    assert all(rcm["risks"][0]["controls"][0]["description"] in c["query"] for c in calls)


def test_rcm_update_during_generation_cannot_commit_stale_draft(client, monkeypatch):
    saved, _ = seed(client, with_source=False)
    url = f"/audits/{saved['audit_id']}/test-plan"
    original = generate_steps

    def change_source(**kwargs):
        result = original(**kwargs)
        rcm = saved["rcm"]
        rcm["risks"][0]["decision"] = "rejected"
        client.post(
            "/rcms",
            json={
                "audit_id": saved["audit_id"],
                "scope": rcm["scope"],
                "risks": rcm["risks"],
            },
        ).raise_for_status()
        return result

    monkeypatch.setattr("app.testplans.router.generate_steps", change_source)
    response = client.post(url + "/generate", json={})
    assert response.status_code == 409
    assert client.get(url).json() is None
