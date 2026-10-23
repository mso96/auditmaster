import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

FIXTURES = Path(__file__).parent / "fixtures"


def test_ingest_markdown_pdf_docx_and_delete(client: TestClient) -> None:
    created = _upload(
        client,
        "ornek-satin-alma-proseduru.md",
        (FIXTURES / "ornek-satin-alma-proseduru.md").read_bytes(),
        "prosedur",
        "Satın Alma",
    )
    detail = client.get(f"/documents/{created['id']}")
    assert detail.status_code == 200
    body = detail.json()
    sections = {chunk["section"] for chunk in body["chunks"]}
    assert "Onay yetkisi" in sections
    assert "Teklif alma" in sections
    assert body["chunk_count"] >= 2
    assert body["audit_area"] == "Satın Alma"

    pdf = _upload(
        client,
        "ornek.pdf",
        _minimal_pdf(["Satin alma onayi zorunludur."]),
        "standart",
        None,
    )
    pdf_detail = client.get(f"/documents/{pdf['id']}").json()
    assert pdf_detail["chunks"][0]["page"] == 1
    assert "onayi" in pdf_detail["chunks"][0]["content"]

    docx = _upload(client, "not.docx", _docx_bytes(), "bulgu", "Satın Alma")
    docx_detail = client.get(f"/documents/{docx['id']}").json()
    assert docx_detail["chunks"][0]["section"] == "Onay yetkisi"
    assert "sipariş" in docx_detail["chunks"][0]["content"]

    listed = client.get("/documents")
    assert listed.status_code == 200
    assert len(listed.json()) == 3

    rejected = client.post(
        "/documents",
        data={"doc_type": "prosedur"},
        files={"file": ("tablo.csv", b"a,b", "text/csv")},
    )
    assert rejected.status_code == 400

    deleted = client.delete(f"/documents/{created['id']}")
    assert deleted.status_code == 204
    assert client.get(f"/documents/{created['id']}").status_code == 404


def test_hybrid_search_folds_turkish_diacritics(client: TestClient) -> None:
    _upload(
        client,
        "ornek-satin-alma-proseduru.md",
        (FIXTURES / "ornek-satin-alma-proseduru.md").read_bytes(),
        "prosedur",
        "Satın Alma",
    )
    _upload(
        client,
        "bordro.txt",
        "Bordro ve maaş hesapları insan kaynakları tarafından tutulur.".encode(),
        "prosedur",
        "İnsan Kaynakları",
    )

    found = client.post("/retrieval/search", json={"query": "satin alma muduru", "limit": 5})
    assert found.status_code == 200
    hits = found.json()["hits"]
    assert hits
    assert hits[0]["filename"] == "ornek-satin-alma-proseduru.md"
    assert hits[0]["lexical_rank"] is not None
    assert hits[0]["vector_rank"] is not None
    assert "müdür" in hits[0]["content"] or "mudur" in hits[0]["content"].casefold()

    filtered = client.post(
        "/retrieval/search",
        json={"query": "teklif", "audit_area": "İnsan Kaynakları", "limit": 5},
    )
    assert filtered.status_code == 200
    assert filtered.json()["hits"]
    assert {hit["filename"] for hit in filtered.json()["hits"]} == {"bordro.txt"}


def test_generate_save_and_export_rcm(client: TestClient) -> None:
    for name, doc_type in (
        ("ornek-satin-alma-proseduru.md", "prosedur"),
        ("ornek-satin-alma-bulgusu.txt", "bulgu"),
        ("ornek-satin-alma-standardi.md", "standart"),
    ):
        _upload(
            client,
            name,
            (FIXTURES / name).read_bytes(),
            doc_type,
            "Satın Alma",
        )

    generated = client.post(
        "/rcm/generate",
        json={
            "scope": "Satın alma denetimi",
            "process_notes": "Eşik 100.000 TL.",
            "audit_area": "Satın Alma",
        },
    )
    assert generated.status_code == 200
    body = generated.json()
    assert body["model_mode"] == "fake"
    assert body["passage_count"] >= 3
    categories = {risk["category"] for risk in body["risks"]}
    assert categories == {"Tedarikçi seçimi", "Yetki ve onay", "Görevler ayrılığı"}
    library = " ".join(
        chunk["content"]
        for document in client.get("/documents").json()
        for chunk in client.get(f"/documents/{document['id']}").json()["chunks"]
    )
    for risk in body["risks"]:
        assert risk["citation_status"] == "kaynaklı"
        assert risk["decision"] == "pending"
        for citation in risk["citations"]:
            assert citation["snippet"] in library
            assert citation["document_name"].startswith("ornek-")

    risks = body["risks"]
    risks[0]["decision"] = "accepted"
    risks[1]["description"] = "Denetçi düzenledi: onay yetkisi olmadan sipariş açılması."
    risks[1]["decision"] = "edited"
    risks[2]["decision"] = "rejected"
    saved = client.post(
        "/rcms",
        json={
            "scope": body["scope"],
            "process_notes": body["process_notes"],
            "audit_area": body["audit_area"],
            "passage_count": body["passage_count"],
            "risks": risks,
        },
    )
    assert saved.status_code == 200
    saved_body = saved.json()
    assert saved_body["rcm"]["risks"][0]["decision"] == "accepted"
    assert saved_body["rcm"]["risks"][1]["decision"] == "edited"
    assert "Denetçi düzenledi" in saved_body["rcm"]["risks"][1]["description"]

    audit = client.get(f"/audits/{saved_body['audit_id']}")
    assert audit.status_code == 200
    assert audit.json()["rcm"]["id"] == saved_body["rcm"]["id"]
    listed = client.get("/audits")
    assert listed.json()[0]["risk_count"] == 3

    exported = client.get(f"/rcms/{saved_body['rcm']['id']}/export")
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(exported.content))
    sheet = workbook["RCM"]
    text = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    assert "Satın alma denetimi" in text
    assert "Tedarikçi seçimi" in text or "Yeterli sayıda yazılı teklif" in text
    assert "Nihai karar denetçiye aittir" in text


def test_generate_without_documents_is_uncited(client: TestClient) -> None:
    response = client.post("/rcm/generate", json={"scope": "Satın alma denetimi"})
    assert response.status_code == 200
    risk = response.json()["risks"][0]
    assert risk["citation_status"] == "kaynaksız"
    assert risk["citations"] == []
    assert risk["controls"][0]["citation_status"] == "kaynaksız"


def test_save_drops_invented_snippet(client: TestClient) -> None:
    _upload(
        client,
        "ornek-satin-alma-proseduru.md",
        (FIXTURES / "ornek-satin-alma-proseduru.md").read_bytes(),
        "prosedur",
        "Satın Alma",
    )
    generated = client.post("/rcm/generate", json={"scope": "Satın alma denetimi"}).json()
    risks = generated["risks"]
    invented = "Bu alıntı belgede yoktur ve uydurulmuştur tamamen."
    for risk in risks:
        for citation in risk["citations"]:
            citation["snippet"] = invented
        for control in risk["controls"]:
            for citation in control["citations"]:
                citation["snippet"] = invented
        for test in risk["tests"]:
            for citation in test["citations"]:
                citation["snippet"] = invented
    saved = client.post(
        "/rcms",
        json={
            "scope": generated["scope"],
            "passage_count": generated["passage_count"],
            "risks": risks,
        },
    )
    assert saved.status_code == 200
    for risk in saved.json()["rcm"]["risks"]:
        assert risk["citation_status"] == "kaynaksız"
        assert risk["citations"] == []
        assert risk["controls"][0]["citation_status"] == "kaynaksız"


def _upload(
    client: TestClient,
    filename: str,
    payload: bytes,
    doc_type: str,
    audit_area: str | None,
) -> dict[str, object]:
    data = {"doc_type": doc_type}
    if audit_area is not None:
        data["audit_area"] = audit_area
    response = client.post(
        "/documents",
        data=data,
        files={"file": (filename, payload, "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Onay yetkisi", level=1)
    document.add_paragraph("Satın alma talepleri onaylanmadan sipariş açılamaz.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _minimal_pdf(lines: list[str]) -> bytes:
    stream = (
        "BT /F1 12 Tf 72 740 Td (" + lines[0].replace("(", "\\(").replace(")", "\\)") + ") Tj ET"
    ).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("latin-1"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("latin-1"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode(
            "latin-1"
        )
    )
    return bytes(output)
